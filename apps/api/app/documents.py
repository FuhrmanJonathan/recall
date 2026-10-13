"""Decide how an uploaded file should be read.

  PDF            -> sent to OpenAI as-is
  .txt / .md     -> decoded locally, then text to OpenAI
  .docx          -> text extracted locally, then text to OpenAI
  png/jpg/webp   -> written notes: local OCR model first, then text to OpenAI
"""

import io
from enum import StrEnum
from pathlib import PurePath

PDF_MAGIC = b"%PDF-"


class SourceKind(StrEnum):
    PDF = "pdf"
    TEXT = "text"
    DOCX = "docx"
    IMAGE = "image"


_EXTENSIONS: dict[str, SourceKind] = {
    ".pdf": SourceKind.PDF,
    ".txt": SourceKind.TEXT,
    ".md": SourceKind.TEXT,
    ".docx": SourceKind.DOCX,
    ".png": SourceKind.IMAGE,
    ".jpg": SourceKind.IMAGE,
    ".jpeg": SourceKind.IMAGE,
    ".webp": SourceKind.IMAGE,
}

SUPPORTED_EXTENSIONS = sorted(_EXTENSIONS)


def detect_source_kind(filename: str | None) -> SourceKind | None:
    """Return how to handle a file, or None if its type isn't supported."""
    return _EXTENSIONS.get(PurePath(filename or "").suffix.lower())


def decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("File must be UTF-8 text.") from exc


def extract_docx_text(data: bytes) -> str:
    """Pull the text (paragraphs, then table cells) out of a .docx file."""
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Could not read this Word document.") from exc

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)
